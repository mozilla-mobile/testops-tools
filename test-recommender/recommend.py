#!/usr/bin/env python3
"""
Firefox iOS Test Recommender

Reads a release diff (between two firefox-vX.Y tags), the TestRail catalogue
export, and the section→module mapping YAML; produces a prioritized Markdown
report of tests to run, risks, and exploratory focus areas.

The pipeline is LLM-assisted (Claude Sonnet 4.6 for rerank, Sonnet 5 for
synthesize by default; both overridable via RECOMMEND_MODEL_* env vars).
Falls back to a deterministic ranker + Markdown writer when the API key is
missing or any LLM call fails — the script never fails closed.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import openpyxl
import yaml

from budget_calculator import (
    BudgetDecision,
    ReleaseSignal,
    compute_test_budget,
    detect_release_type,
)
from candidate_scorer import (
    ScoringContext,
    build_scoring_context,
    pre_filter_candidates,
)
from nimbus_flag_analyzer import (
    FeatureFlag,
    attribute_files,
    classify_gating_flags,
    load_flag_catalogue,
    parse_wrapper_mapping,
)

try:
    import anthropic
    _ANTHROPIC_AVAILABLE = True
except ImportError:
    _ANTHROPIC_AVAILABLE = False

REPO = "mozilla-mobile/firefox-ios"

# Two-model split, decided empirically after A/B testing four configs
# (see metrics_baseline.md and Phase 1 comparison notes):
#   - Rerank on Sonnet 4.6 with temperature=0 gave the best output stability
#     (48% Jaccard vs 28% for Haiku, 39% for Sonnet 5 default).
#   - Synthesize on Sonnet 5 gives better prose than 4.6 at similar cost.
# Both overridable via env vars for quick rollback if a model change regresses
# output quality on a future release.
LLM_MODEL_RERANK = os.environ.get("RECOMMEND_MODEL_RERANK", "claude-sonnet-4-6")
LLM_MODEL_SYNTHESIZE = os.environ.get("RECOMMEND_MODEL_SYNTHESIZE", "claude-sonnet-5")

# Paths that are noise for "module touched" classification: build/project files,
# localization, assets, and the test trees themselves (the test tree is used for
# exact-match against TestRail Automated Test Name, but it is NOT a "product
# module" we'd ask QA to retest).
NOISE_PATH_PATTERNS = (
    re.compile(r"\.xcodeproj(/|$)"),
    re.compile(r"\.xcworkspace(/|$)"),
    re.compile(r"Info\.plist$"),
    re.compile(r"\.entitlements$"),
    re.compile(r"\.lproj/"),
    re.compile(r"/Assets\.xcassets/"),
    re.compile(r"(^|/)Assets/"),
    re.compile(r"Localizable\.strings$"),
    re.compile(r"\.strings$"),
    re.compile(r"\.gitignore$"),
    re.compile(r"\.md$"),
    re.compile(r"^\.github/"),
)

TEST_PATH_PATTERNS = (
    re.compile(r"/Tests/"),
    re.compile(r"Tests\.swift$"),
    re.compile(r"-tests/"),
)


def is_noise_path(p: str) -> bool:
    return any(rx.search(p) for rx in NOISE_PATH_PATTERNS)


def is_test_path(p: str) -> bool:
    return any(rx.search(p) for rx in TEST_PATH_PATTERNS)


# =============================================================================
# Data models
# =============================================================================


@dataclass
class FileChange:
    path: str
    additions: int
    deletions: int
    patch: str = ""       # truncated diff text (for grep heuristics + LLM)


@dataclass
class PR:
    number: int
    title: str
    author: str
    additions: int
    deletions: int


@dataclass
class TestCase:
    id: str
    title: str
    section_top: str           # top-level section, e.g. "Library"
    section_hierarchy: str     # full path
    sub_suite: str
    automation: str
    automated_test_name: Optional[str]


@dataclass
class RiskSignal:
    kind: str
    severity: str              # "high" | "medium" | "low"
    location: str              # "PR #1234" or "BrowserKit/Sources/WebEngine/Foo.swift"
    detail: str


@dataclass
class ModuleChange:
    module: str
    files: list[FileChange] = field(default_factory=list)

    @property
    def total_loc(self) -> int:
        return sum(f.additions + f.deletions for f in self.files)


@dataclass
class DriftFinding:
    kind: str                  # "testrail_new_section", "testrail_stale_section",
                               # "repo_new_module", "repo_dead_path"
    item: str
    detail: str


@dataclass
class Analysis:
    """The full deterministic analysis, before LLM rerank/synthesize."""
    from_tag: str
    to_tag: str
    prs: list[PR]
    skipped_prs: list[tuple[PR, str]]
    file_changes: list[FileChange]
    module_changes: dict[str, ModuleChange]
    risks: list[RiskSignal]
    drift: list[DriftFinding]
    exact_matched_tests: list[TestCase]          # by Automated Test Name path
    section_matched_tests: list[TestCase]        # by mapping.yaml lookup
    unclassified_files: list[FileChange]
    budget: BudgetDecision                       # test count budget (Phase 2)
    scoring_context: ScoringContext              # for pre-filter (Phase 3)
    # Nimbus flag analysis. Empty catalogue = --repo-path not provided; the
    # pipeline degrades gracefully to the pre-Nimbus behavior.
    flag_catalogue: dict[str, FeatureFlag] = field(default_factory=dict)
    module_flag_states: dict[str, dict] = field(default_factory=dict)
    # section_top → classification ("unflagged" | "enabled" | "beta_only" | "fully_dark")
    section_flag_state: dict[str, str] = field(default_factory=dict)


# =============================================================================
# Loaders
# =============================================================================


def load_mapping(path: Path) -> dict:
    with path.open() as f:
        return yaml.safe_load(f)


def load_testrail(path: Path) -> list[TestCase]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # TestRail exports have shipped with different default sheet names
    # over the years ("Worksheet", "Sheet1", localized variants…). Use the
    # first sheet regardless of its name.
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    headers = list(rows[0])
    idx = {h: i for i, h in enumerate(headers)}

    cases: list[TestCase] = []
    for r in rows[1:]:
        sh = r[idx["Section Hierarchy"]] or ""
        cases.append(TestCase(
            id=str(r[idx["ID"]] or ""),
            title=str(r[idx["Title"]] or ""),
            section_top=sh.split(">")[0].strip() if sh else "",
            section_hierarchy=sh,
            sub_suite=str(r[idx["Sub Test Suite(s)"]] or ""),
            automation=str(r[idx["Automation"]] or ""),
            automated_test_name=(str(r[idx["Automated Test Name(s)"]]) if r[idx["Automated Test Name(s)"]] else None),
        ))
    return cases


# =============================================================================
# GitHub diff fetching (via gh CLI)
# =============================================================================


def gh_json(args: list[str]) -> dict | list:
    """Run a gh api call and return parsed JSON. Exits on failure."""
    cmd = ["gh", "api"] + args
    out = subprocess.run(cmd, capture_output=True, text=True)
    if out.returncode != 0:
        sys.stderr.write(f"gh api failed: {' '.join(cmd)}\n{out.stderr}\n")
        sys.exit(1)
    return json.loads(out.stdout)


def fetch_compare(from_tag: str, to_tag: str, max_files: int = 300) -> tuple[list[FileChange], list[dict]]:
    """Return (file_changes, commits) between two tags via the GitHub compare API.

    The compare endpoint caps its `files` array at 300 entries and does NOT
    paginate them (page=2 comes back empty), so a bigger diff cannot be fetched
    this way at all. We abort loudly rather than truncate, and point at the
    git-first path (--repo-path), which has no cap.
    """
    data = gh_json([f"repos/{REPO}/compare/{from_tag}...{to_tag}?per_page=300"])
    files = []
    for f in data.get("files", [])[:max_files]:
        files.append(FileChange(
            path=f["filename"],
            additions=f.get("additions", 0),
            deletions=f.get("deletions", 0),
            patch=(f.get("patch") or "")[:8000],   # truncate per-file patch
        ))
    if len(data.get("files", [])) >= max_files:
        sys.stderr.write(
            f"ERROR: diff has >= {max_files} files (GitHub compare API cap; "
            f"the report would silently omit changes). Aborting.\n"
            f"Re-run with --repo-path <local firefox-ios clone> to compute the "
            f"diff from local git instead (no cap).\n"
        )
        sys.exit(2)
    return files, data.get("commits", [])


# -----------------------------------------------------------------------------
# Git-first diff (no 300-file cap) — used when --repo-path is provided
# -----------------------------------------------------------------------------


def git_out(repo_path: Path, args: list[str]) -> str:
    """Run a git command in repo_path and return stdout. Exits on failure."""
    cmd = ["git", "-C", str(repo_path), "-c", "core.quotepath=false"] + args
    out = subprocess.run(cmd, capture_output=True, text=True)
    if out.returncode != 0:
        sys.stderr.write(f"git failed: {' '.join(cmd)}\n{out.stderr}\n")
        sys.exit(1)
    return out.stdout


def resolve_ref(repo_path: Path, ref: str) -> str:
    """Resolve a tag/branch/SHA to a commit, trying origin/<ref> as a fallback
    so callers can pass `release/v154.0` without the remote prefix."""
    for candidate in (ref, f"origin/{ref}"):
        out = subprocess.run(
            ["git", "-C", str(repo_path), "rev-parse", "--verify", "-q", f"{candidate}^{{commit}}"],
            capture_output=True, text=True,
        )
        if out.returncode == 0:
            return out.stdout.strip()
    sys.stderr.write(
        f"ERROR: ref '{ref}' not found in {repo_path}. "
        f"Try: git -C {repo_path} fetch origin --tags --prune\n"
    )
    sys.exit(2)


def parse_numstat_z(raw: str) -> dict[str, tuple[int, int]]:
    """Parse `git diff -M --numstat -z` into {path: (additions, deletions)}.

    In -z form a normal entry is "adds\\tdels\\tpath\\0"; a rename has an empty
    path field followed by two extra NUL-terminated records (old, new). Binary
    files report "-" for both counts.
    """
    parts = raw.split("\0")
    stats: dict[str, tuple[int, int]] = {}
    i = 0
    while i < len(parts):
        rec = parts[i]
        i += 1
        if not rec:
            continue
        fields = rec.split("\t")
        if len(fields) < 3:
            continue
        adds, dels, path = fields[0], fields[1], fields[2]
        if path == "":                      # rename/copy: old and new follow
            if i + 1 >= len(parts):
                break
            path = parts[i + 1]             # keep the new path
            i += 2
        a = int(adds) if adds.isdigit() else 0
        d = int(dels) if dels.isdigit() else 0
        stats[path] = (a, d)
    return stats


def parse_unified_diff(raw: str) -> dict[str, str]:
    """Split `git diff` output into {path: patch}, where patch starts at the
    first @@ hunk header — same shape as the GitHub compare API's `patch`."""
    patches: dict[str, str] = {}
    path: Optional[str] = None
    hunks: list[str] = []
    in_header = False                       # between "diff --git" and the first @@

    def flush() -> None:
        if path and hunks:
            patches[path] = "\n".join(hunks)

    for line in raw.split("\n"):
        if line.startswith("diff --git "):
            flush()
            path, hunks, in_header = None, [], True
        elif in_header and line.startswith("+++ b/"):
            # git appends a TAB after the name when the path contains spaces.
            path = line[6:].split("\t")[0]
        elif in_header and line.startswith("--- a/") and path is None:
            path = line[6:].split("\t")[0]  # deletion: +++ is /dev/null
        elif in_header and line.startswith("@@"):
            # First hunk header. Everything from here is patch body — including
            # lines that look like diff headers (diffs of diffs, .patch files).
            in_header = False
            hunks.append(line)
        elif not in_header and path is not None:
            hunks.append(line)
    flush()
    return patches


def fetch_compare_git(repo_path: Path, from_ref: str, to_ref: str) -> tuple[list[FileChange], list[dict]]:
    """Return (file_changes, commits) computed from a local clone.

    Uses merge-base..to (three-dot semantics) to match what the GitHub compare
    API reports, but without its 300-file cap.
    """
    base_commit = resolve_ref(repo_path, from_ref)
    head_commit = resolve_ref(repo_path, to_ref)
    merge_base = git_out(repo_path, ["merge-base", base_commit, head_commit]).strip()
    rng = f"{merge_base}..{head_commit}"

    stats = parse_numstat_z(git_out(repo_path, ["diff", "-M", "--numstat", "-z", rng]))
    patches = parse_unified_diff(git_out(repo_path, ["diff", "-M", "-U3", rng]))

    files = [
        FileChange(
            path=path,
            additions=adds,
            deletions=dels,
            patch=patches.get(path, "")[:8000],   # truncate per-file patch
        )
        for path, (adds, dels) in stats.items()
    ]
    commits = [{"sha": sha} for sha in git_out(repo_path, ["log", "--format=%H", rng]).split()]
    return files, commits


def fetch_prs_for_commits(commits: list[dict], limit: int = 500) -> list[PR]:
    """For each commit, find its PR and then fetch the full PR for accurate
    additions/deletions (the /commits/<sha>/pulls payload omits those).

    TODO(git-first): replace with git_pr_extractor.build_prs_from_git — the
    module is fully tested but not yet integrated (see NEXT_STEPS.md
    Priority 1). Removing this function eliminates ~281 API calls per major.
    """
    pr_numbers: set[int] = set()
    pr_basics: dict[int, dict] = {}
    for c in commits[:limit]:
        sha = c["sha"]
        try:
            data = gh_json([f"repos/{REPO}/commits/{sha}/pulls"])
        except SystemExit:
            continue
        for p in data:
            n = p["number"]
            if n not in pr_numbers:
                pr_numbers.add(n)
                pr_basics[n] = p

    prs: list[PR] = []
    for n in sorted(pr_numbers):
        try:
            full = gh_json([f"repos/{REPO}/pulls/{n}"])
        except SystemExit:
            full = pr_basics[n]
        prs.append(PR(
            number=n,
            title=full.get("title", ""),
            author=(full.get("user") or {}).get("login", ""),
            additions=full.get("additions", 0),
            deletions=full.get("deletions", 0),
        ))
    return prs


# =============================================================================
# PR filtering (low-impact)
# =============================================================================


LOW_IMPACT_TITLE_PREFIXES = (
    "string import",
    "strings import",
    "[v",                    # version bumps in title
    "bump ",
)

LOW_IMPACT_TITLE_KEYWORDS = (
    "string import",
    "strings import",
    "l10n",
    "localization",
)


def is_low_impact_pr(pr: PR) -> Optional[str]:
    """Return reason string if this PR should be filtered out, else None."""
    title_lc = pr.title.lower()
    if any(title_lc.startswith(p) for p in LOW_IMPACT_TITLE_PREFIXES):
        return f"title prefix: {pr.title[:60]}"
    if any(k in title_lc for k in LOW_IMPACT_TITLE_KEYWORDS):
        return f"localization/strings: {pr.title[:60]}"
    # Mergify backport detection happens at branch level (not visible from PR
    # number alone); we approximate via title.
    if "[backport]" in title_lc or "mergify" in title_lc:
        return "mergify backport"
    return None


# =============================================================================
# Risk heuristics (cheap, no LLM)
# =============================================================================


CONCURRENCY_RE = re.compile(r"\b(async|await|actor|DispatchQueue|Task\.|withCheckedContinuation|@MainActor|@Sendable)\b")
FORCE_UNWRAP_RE = re.compile(r"(?<![A-Za-z_])try!|fatalError|preconditionFailure|!\s*\.|\![\s\)\.,;]")
ERROR_HANDLING_RE = re.compile(r"\b(throw|throws|try\b|catch|Result<)")

DEPENDENCY_PATHS = (
    "Package.swift",
    "Package.resolved",
    "Podfile",
    "Podfile.lock",
    "MozillaRustComponents/",
)

# TODO(nimbus-fml): replace this coarse "any change to Nimbus config = high risk"
# heuristic with FML-based auto-detection. Currently every touched file under
# nimbus-features/ produces a "high" severity risk regardless of whether the
# feature is off_in_release. Planned: parse the FML defaults per channel
# between --from and --to tags, only flag as risk when the release-channel
# `enabled` state changed, and route off-in-release features to a
# "Nightly-only tests" report section. Note: fewer 'high' severity risks
# means budget_calculator will bump less often — expected side effect.
NIMBUS_PATHS = (
    "firefox-ios/nimbus.fml.yaml",
    "firefox-ios/nimbus-features/",
    "firefox-ios/initial_experiments.json",
)


def patch_added_lines(patch: str) -> str:
    """Return only the lines added by the patch (lines starting with +, excluding +++)."""
    out = []
    for line in patch.splitlines():
        if line.startswith("+") and not line.startswith("+++"):
            out.append(line[1:])
    return "\n".join(out)


def detect_risks(prs: list[PR], file_changes: list[FileChange]) -> list[RiskSignal]:
    """Compute deterministic risk signals from PR-level and file-level heuristics.

    NOTE: RiskSignal.location has an undocumented-but-load-bearing contract
    used by candidate_scorer.build_scoring_context. It must be either:
      - "PR #N"           → PR-level, skipped by the scorer's section mapping
      - a file path with "/" → file-level, matched against touched modules
    Locations without "/" (e.g. bare "Package.swift", "Podfile.lock") are
    intentionally excluded from the "sections with risk" score boost, since
    they don't map to any single TestRail section. If you add a new risk
    type whose location doesn't fit either shape, either fix its location
    or update the scorer's filter (candidate_scorer.py ~line 92).
    """
    risks: list[RiskSignal] = []

    # Per-PR signals
    for pr in prs:
        loc = pr.additions + pr.deletions
        if loc > 1000:
            risks.append(RiskSignal("large_pr", "high", f"PR #{pr.number}",
                                    f"{loc} LOC, '{pr.title[:80]}' by @{pr.author}"))

    # Per-file content signals (uses patch text from compare endpoint).
    # Test and noise files are excluded — these flags are about product risk,
    # not test code style.
    for fc in file_changes:
        if not fc.patch:
            continue
        if is_noise_path(fc.path) or is_test_path(fc.path) or "Mock" in fc.path:
            continue
        added = patch_added_lines(fc.patch)
        if CONCURRENCY_RE.search(added):
            risks.append(RiskSignal("concurrency", "medium", fc.path,
                                    "added/modified concurrent code (async/await/actor/DispatchQueue)"))
        if FORCE_UNWRAP_RE.search(added):
            risks.append(RiskSignal("force_unwrap", "low", fc.path,
                                    "added force-unwrap, try!, or fatalError"))
        if ERROR_HANDLING_RE.search(added) and ".swift" in fc.path:
            if fc.additions > 50:
                risks.append(RiskSignal("error_handling", "low", fc.path,
                                        f"error-handling changes ({fc.additions} added lines)"))

    # Release-level signals (deps + nimbus)
    touched_paths = {fc.path for fc in file_changes}
    for dep in DEPENDENCY_PATHS:
        if any(p.startswith(dep) or p == dep for p in touched_paths):
            hits = [p for p in touched_paths if p.startswith(dep) or p == dep]
            risks.append(RiskSignal("dependency", "medium", dep,
                                    f"dependency files changed: {', '.join(hits[:3])}"))
    for n in NIMBUS_PATHS:
        if any(p.startswith(n) for p in touched_paths):
            hits = [p for p in touched_paths if p.startswith(n)]
            risks.append(RiskSignal("nimbus_flags", "high", n,
                                    f"Nimbus/feature-flag config changed — behavior may flip server-side without code change in product. Files: {', '.join(hits[:3])}"))

    return risks


# =============================================================================
# Module classification
# =============================================================================


def known_modules_from_mapping(mapping: dict) -> list[str]:
    """Collect every code path referenced anywhere in the mapping."""
    paths: set[str] = set()
    for s in mapping.get("sections", []):
        for m in s.get("modules", []):
            paths.add(m["path"])
    for p in mapping.get("modules_without_clear_section", []) or []:
        paths.add(p)
    return sorted(paths, key=len, reverse=True)   # longest first for longest-prefix match


def classify_file(path: str, known: list[str]) -> Optional[str]:
    """Map a changed file to a known module path (longest-prefix match)."""
    for m in known:
        if path == m or path.startswith(m.rstrip("/") + "/"):
            return m
    return None


def group_by_module(file_changes: list[FileChange], prs: list[PR], mapping: dict) -> tuple[dict[str, ModuleChange], list[FileChange]]:
    """Group changed product-code files by module path. Test files and noise
    paths are excluded — tests are handled separately via exact match, noise
    is dropped entirely. Return (groups, unclassified_product_files)."""
    known = known_modules_from_mapping(mapping)
    groups: dict[str, ModuleChange] = {}
    unclassified: list[FileChange] = []

    for fc in file_changes:
        if is_noise_path(fc.path) or is_test_path(fc.path):
            continue
        m = classify_file(fc.path, known)
        if m is None:
            unclassified.append(fc)
            continue
        if m not in groups:
            groups[m] = ModuleChange(module=m)
        groups[m].files.append(fc)
    return groups, unclassified


# =============================================================================
# Test matching
# =============================================================================


def normalize_automated_path(s: str) -> list[str]:
    """Extract Swift file paths from an Automated Test Name(s) field, which may
    contain multiple entries separated by newlines, with formats like:
        firefox-ios/firefox-ios-tests/Tests/XCUITests/TodayWidgetTests.swift#testFoo()
        Tests/XCUITests/ModernKitOnboardingTests/testFoo
        XCUITests/IntegrationTests#testFoo
    Return the file-path portion(s) without the # suffix.
    """
    paths = []
    for entry in re.split(r"[\n,]+", s):
        entry = entry.strip()
        if not entry:
            continue
        path = entry.split("#", 1)[0].strip()
        if path:
            paths.append(path)
    return paths


def exact_match_by_test_file(file_changes: list[FileChange], tests: list[TestCase]) -> list[TestCase]:
    """Find TestRail cases whose Automated Test Name references a changed file."""
    changed_test_files = [fc.path for fc in file_changes if "/Tests/" in fc.path or fc.path.endswith("Tests.swift")]
    matched: list[TestCase] = []
    for tc in tests:
        if not tc.automated_test_name:
            continue
        for ref in normalize_automated_path(tc.automated_test_name):
            # match if any changed test file ends-with or contains this ref tail
            tail = ref.split("Tests/", 1)[-1]
            for cf in changed_test_files:
                if tail and tail in cf:
                    matched.append(tc)
                    break
            else:
                continue
            break
    return matched


def section_match_tests(modules_touched: list[str], mapping: dict, tests: list[TestCase]) -> list[TestCase]:
    """For each touched module, find sections that map to it, then return tests in those sections."""
    module_to_sections: dict[str, set[str]] = defaultdict(set)
    for s in mapping.get("sections", []):
        name = s["name"]
        for m in s.get("modules", []):
            module_to_sections[m["path"]].add(name)

    relevant_sections: set[str] = set()
    for mt in modules_touched:
        # exact + parent matches (since modules in YAML may be coarser than touched paths)
        for known_module, sections in module_to_sections.items():
            if mt == known_module or mt.startswith(known_module.rstrip("/") + "/") or known_module.startswith(mt.rstrip("/") + "/"):
                relevant_sections.update(sections)

    return [tc for tc in tests if tc.section_top in relevant_sections]


# =============================================================================
# Drift detection
# =============================================================================


def detect_drift(file_changes: list[FileChange], tests: list[TestCase], mapping: dict) -> list[DriftFinding]:
    findings: list[DriftFinding] = []

    # TestRail-side: sections in export but not in YAML
    yaml_sections = {s["name"] for s in mapping.get("sections", [])}
    export_sections = {tc.section_top for tc in tests if tc.section_top}
    new_in_export = export_sections - yaml_sections
    stale_in_yaml = yaml_sections - export_sections
    for s in sorted(new_in_export):
        sample_titles = [tc.title for tc in tests if tc.section_top == s][:5]
        findings.append(DriftFinding("testrail_new_section", s,
                                     f"sample test titles: {sample_titles}"))
    for s in sorted(stale_in_yaml):
        findings.append(DriftFinding("testrail_stale_section", s,
                                     "section name in YAML not present in latest export (rename or removal?)"))

    # Repo-side: touched product-code paths under known parent dirs but not in any YAML module.
    # Test files and noise (Info.plist, .xcodeproj, .lproj/, Assets/) are excluded.
    known = known_modules_from_mapping(mapping)
    parents = ("BrowserKit/Sources/", "firefox-ios/Client/Frontend/", "firefox-ios/Client/", "firefox-ios/")
    unclassified_modules: set[str] = set()
    for fc in file_changes:
        if is_noise_path(fc.path) or is_test_path(fc.path):
            continue
        if classify_file(fc.path, known) is not None:
            continue
        for parent in parents:
            if fc.path.startswith(parent):
                rest = fc.path[len(parent):]
                head = rest.split("/", 1)[0]
                if head:
                    unclassified_modules.add(parent + head)
                break
    for m in sorted(unclassified_modules):
        findings.append(DriftFinding("repo_new_module", m,
                                     "module appears in diff but is not referenced in mapping YAML"))

    return findings


# =============================================================================
# LLM rerank + synthesize  (models configured via LLM_MODEL_RERANK / LLM_MODEL_SYNTHESIZE)
# =============================================================================
# Design notes:
#   - Two stages: rerank produces a prioritized subset with structured output;
#     synthesize produces the Markdown narrative.
#   - Pool sizing: candidates are pre-filtered deterministically (see
#     candidate_scorer.py) to top_k = budget.final_hi * 4 BEFORE reaching the
#     LLM. Typical minor release: ~989 raw → ~280 after pre-filter.
#   - Prompt caching: both stages put stable instructions in the `system`
#     block with cache_control:ephemeral. The 5-min TTL means cross-release
#     cache misses (releases are days apart) but same-session reruns hit
#     cache — useful for retries, dry-runs, and back-to-back experiments.
#   - Graceful fallback: if ANTHROPIC_API_KEY is unset or the SDK isn't
#     installed, or any API call fails, fall back to the deterministic ranker
#     (`_deterministic_rerank`) and the deterministic Markdown writer
#     (`render_deterministic_report`). The pipeline never fails closed.
#   - Determinism: temperature=0 on Sonnet 4.6 (Sonnet 5+ deprecated
#     temperature; effort=low is the equivalent). Combined with the
#     deterministic pre-filter, run-to-run overlap on identical inputs is
#     ~55% Jaccard on our v151.2→v151.3 baseline. Residual variance is cloud
#     non-determinism; not reducible without ensemble or LLM bypass.
#   - Effort: "low" for rerank (classification-shaped task), "medium" for
#     synthesize (writing a structured report). Skipped for Haiku (rejects).
#   - Thinking: disabled — neither stage needs multi-step reasoning, and
#     disabling avoids spending tokens on it.


@dataclass
class RankedTest:
    test_id: str
    priority: str           # "P0" | "P1" | "P2"
    reason: str


SYSTEM_PROMPT_RERANK = """You are a release QA test-prioritization assistant for the Firefox iOS application (github.com/mozilla-mobile/firefox-ios). The manual QA team cannot execute the full ~1650-case regression suite on every release, so you select the highest-value subset to run.

# Your task
Given:
  1. A list of "candidate" TestRail test cases (already filtered to areas plausibly affected by this release). Each candidate includes a deterministic `score` field (integer, can be negative). Higher score means stronger a-priori signal: exact code↔test match, section maps to a high-LOC touched module, section has an attached risk signal, or the test is manual-only. Negative scores are for tests already covered by CI. Use `score` as a strong hint but not a hard constraint — you may still promote a lower-score candidate if the risk context justifies it, or demote a high-score candidate if the reason is thin.
  2. A summary of the modules touched in the release diff
  3. A list of risk signals computed from the diff (concurrency changes, dependency bumps, Nimbus feature-flag changes, force-unwraps, hotspots, large PRs without tests, etc.)
  4. A list of "drift" findings (modules touched that are not yet mapped to any TestRail section — coverage blind spots)

Return a prioritized subset of {budget_lo}-{budget_hi} tests, each tagged with:
  - priority: "P0" | "P1" | "P2"
  - reason: one short sentence (under 25 words) explaining why this test is worth running for THIS release specifically (not generic descriptions of the test)

# Priority rubric

P0 (must run)
  - Directly tests behavior in a module that was modified with substantial LOC (>100) AND has risk signals attached, OR
  - Covers a flow that depends on a changed dependency (Rust components, Package.swift), OR
  - Cross-references a Nimbus feature flag that changed, OR
  - Tests a critical user flow (authentication, sync, tab management, search) where the touched module is implicated.

P1 (should run)
  - Tests behavior in a touched module but with lower change volume or no risk signals, OR
  - Adjacent to a P0 test in the same section (smoke-coverage of the area).

P2 (run if time)
  - Tests in sections only loosely connected to the changes (e.g. shared UI components, settings adjacent to changed areas).
  - DO NOT include tests purely because they share a section with a touched module if no real connection exists.

# Important rules
  - Tests marked `automation: Completed` are already covered by CI — DEPRIORITIZE these (P2 at best) unless the change directly modified the underlying tests too.
  - Smoke & Sanity tests are handled separately by the automated smoke suite and are NOT in the candidate list. Do not ask for them; focus on the release-specific regression the candidates represent.
  - Tests marked `automation: Unsuitable` are manual-only and must be run by humans — favor these for P0/P1 when relevant.
  - Feature-flag rule (channel-aware):
      · `flag_state: "fully_dark"` — the candidate exercises code gated by a Nimbus flag that is OFF in release, beta AND developer. Real users cannot reach this code. Force to P2 and the `reason` must acknowledge the flag state (e.g. "code shipped but flag OFF in all channels; verify only after manual flag toggle").
      · `flag_state: "beta_only"` — OFF in release, ON in beta or developer. Dogfooding users see it but no release users. P1 at most; the `reason` should note "dogfooding-only for this release".
      · `flag_state: "enabled"` — ON in release. Treat normally.
      · No `flag_state` field — no flag detected. Treat normally.
    This overrides the rest of the priority rubric — a fully_dark candidate never gets P0/P1 even if it would otherwise qualify.
  - Do not invent test IDs. Only return IDs present in the candidate list.
  - Hard cap: return at most {budget_hi} ranked tests.
  - Be honest about uncertainty in the `reason` field. A reason like "tests the touched module" is fine if you genuinely don't have more signal.

# Output format
You will respond using the provided JSON schema. Order the `ranked_tests` array by priority (P0 first, then P1, then P2). Within the same priority, order by the section that received the most changes first.

In `notes`, briefly list (under 200 words):
  - Any coverage gaps you noticed (modules touched but with no candidates that meaningfully cover them)
  - Any flows where you think a new automated regression test would be valuable
  - Any areas suitable for exploratory testing
"""


SYSTEM_PROMPT_SYNTHESIZE = """You are a release QA report writer for the Firefox iOS application. You produce a single Markdown report per release that the manual QA team uses to decide what to test.

# Inputs you receive
  1. The release tag pair (from → to)
  2. Module-level summary of code changes
  3. Risk signals (concurrency, force-unwraps, dependency bumps, Nimbus changes, large PRs, hotspots)
  4. Mapping drift findings (modules touched without TestRail coverage — these are coverage blind spots)
  5. A ranked list of TestRail tests with priorities (P0/P1/P2) and per-test reasons, plus narrative notes from the rerank stage
  6. Unclassified file changes (touched paths the pipeline could not map to any module)

# Report structure (use exactly these section headings, in this order)

## Executive summary
2-4 bullets. What is the headline of this release? Top 2-3 areas the QA team should focus on. Mention the largest single risk if one stands out.
CRITICAL: when calling out "new features shipping", check `flag_state` on the corresponding modules FIRST. A module with `flag_state: "fully_dark"` is code that ships in the binary but is NOT reachable by any user (release / beta / developer) — do NOT frame it as "ships this release". Frame it as "landed behind an OFF Nimbus flag — no user exposure until rollout starts". A `flag_state: "beta_only"` module reaches dogfooding channels only — say "landed behind flag `<accessor>` — currently ON in beta/developer only, OFF in release".

## Suggested manual tests
Three subsections: ### P0 (must run), ### P1 (should run), ### P2 (if time).
Under each, list the tests like:
  - `C123456` — Test title — *Reason: ...*  _(section, sub_suite, automation status)_
Group by section_top when there are 5+ tests in the same area.

## Coverage gaps — manual tests to add
Identify areas where:
  - Modules were touched but no TestRail tests exist for them (use the drift findings + unclassified files)
  - Changes introduce behavior that the existing catalog plausibly does not cover
  - The QA team should consider adding a new manual test case
Be specific. Cite the touched module or PR pattern.

## Suggested automated regression
For critical flows touched in this release that lack automated regression:
  - Identify candidate flows (auth, sync, tabs, search, etc.) where the diff suggests new behavior or risk
  - Suggest where a new automated regression test would prevent future drift
  - Prefer suggestions where `automation: Unsuitable` is high — those are manual-only and the team is bottlenecked there

## Feature flag status
CRITICAL — how to reason about flag state:
1. `flag_catalogue` (top-level) is the AUTHORITATIVE list of every Nimbus flag in the repo with its per-channel state. If you can plausibly identify a touched module with an accessor in this catalogue by name similarity (compare tokens in `module_changes[].module` and `module_changes[].files` against the `accessor` field in `flag_catalogue`), USE the catalogue's state — never guess from world knowledge or general feature descriptions.
2. `module_changes[].flag_state` is a stronger signal: the pipeline detected a direct code reference to a flag in that module. When present, prefer it over any name-based inference.
3. If neither per-module attribution nor a plausible catalogue name-match supports a flag-state claim, DO NOT invent one. Write "flag state not detected — verify manually before assuming release exposure" and move on. Never say "enabled in release" or "user-facing" as a default when you have no evidence.
4. The text grep is imperfect and misses indirect gating (dependency injection, call chains). Absence of `flag_state` on a module is NOT evidence of "always on" — it just means our detector didn't find a reference. Rely on the catalogue name-match to fill this gap.

List every module in `module_changes` where `flag_state` is present, grouped by classification. Format each line as:
  - `<module_path>` — flag `<accessor>` (`<release_state>`) — <what this means for QA>

Only include this section when there is at least one gated module in this release. If all touched modules are unflagged, write "No feature-gated changes in this release" as a one-liner and move on.

The three classifications and their framing:
  - **fully_dark**: OFF in release, beta, and developer. Not reachable by anyone. QA action: verify only after a manual flag toggle in a debug build; do NOT allocate release-regression time.
  - **beta_only**: OFF in release, ON in beta/developer. Reachable by dogfooding users. QA action: exercise via the beta build; not a regression risk for release users.
  - **enabled**: ON in release. Reachable by all users. QA action: standard release regression coverage applies.

Cite the accessor name and the compact state summary from `gating_flags[].state` so the reader knows exactly which flag and its per-channel setting.

## Risks
Translate the risk signals into prose, ranked high → low severity:
  - Concurrency changes — name the files, explain why they matter
  - Force-unwrap / try! introductions — risk of runtime crashes
  - Dependency bumps (especially MozillaRustComponents) — explain downstream impact (Sync, Logins, Autofill)
  - Nimbus / feature-flag changes — flag that production behavior may differ from build behavior depending on rollout
  - Large PRs (>1000 LOC) — name the PR titles, suggest reviewers
  - Hotspot files (modified by many authors recently) — flag for closer review

## Exploratory testing focus
Suggest 3-6 specific areas of the app for exploratory testing this release, framed as "Spend 20-30 min poking at X with focus on Y". Tie each suggestion to a concrete change in the diff.

## High-risk PRs to review more deeply
A short list of PRs that warrant deeper code review beyond QA — large refactors, hotspot-heavy changes, dependency bumps, or PRs that introduced multiple risk signals. One bullet per PR with the # and a one-line justification.

# Examples of good vs bad output

For the Risks section, prefer specific and actionable framings over generic warnings.

BAD:  "PR #33846 introduces concurrency changes that may cause issues."
GOOD: "PR #33846 — ReaderModeSchemeHandler refactor (1142 LOC, @Alex-Bangu): async/await plus a new force-unwrap in the scheme handler that every Reader Mode entry point depends on. A crash here is user-visible on every article page."

For Coverage gaps, name the exact files or modules and propose specific test scenarios rather than restating that coverage is missing.

BAD:  "Add tests for the WorldCup widget."
GOOD: "WorldCup widget (989 LOC across 7 files, including WorldCupWinnerBackgroundView.swift with a force-unwrap) — no TestRail cases exist. Add: (1) widget renders knockout-phase brackets correctly; (2) widget handles API error state without crash; (3) widget survives homepage foreground-resume without duplicating requests."

For High-risk PRs, keep bullets to one line each. The reader wants to skim, not to read a paragraph per PR.

GOOD: "- #33846 (Reader mode scheme refactor, 1142 LOC): async concurrency plus a new force-unwrap on the critical Reader Mode path — warrants senior iOS review."

Anti-patterns to avoid in every section:
  - Hedge language ("could potentially", "may possibly", "it is worth noting"). Lead with the fact.
  - Rewriting the input JSON as prose. The reader has the diff; give them insight, not restatement.
  - Bullets that end in "…" or trail off. Every bullet must land a claim.

# Style guide
  - English, professional, direct.
  - No filler ("It is important to note that..."). Lead with the fact.
  - Use Markdown lists liberally; minimize prose paragraphs.
  - Cite test IDs as backticked codes: `C123456`.
  - Cite paths as backticked: `firefox-ios/Client/Frontend/Reader/...`.
  - Never invent details. If a section has no content, write a one-line "No findings in this release" instead of fabricating.
  - The report goes to a QA lead who will read it in 5 minutes — every word should earn its place.
"""


RERANK_SCHEMA = {
    "type": "object",
    "properties": {
        "ranked_tests": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "test_id": {"type": "string"},
                    "priority": {"type": "string", "enum": ["P0", "P1", "P2"]},
                    "reason": {"type": "string"},
                },
                "required": ["test_id", "priority", "reason"],
                "additionalProperties": False,
            },
        },
        "notes": {"type": "string"},
    },
    "required": ["ranked_tests", "notes"],
    "additionalProperties": False,
}


def _llm_available() -> bool:
    return _ANTHROPIC_AVAILABLE and bool(os.environ.get("ANTHROPIC_API_KEY"))


def _candidates_for_prompt(analysis: Analysis) -> tuple[list[TestCase], list[dict]]:
    """Dedup exact + section matches, pre-filter to top-K by deterministic
    score, and return (TestCase list, compact dicts for prompt).

    Two-stage narrowing:
      1. Dedup + drop Smoke & Sanity (Smoke runs automatically).
      2. Score each remaining candidate; keep only the top-K where
         K = budget.final_hi * 4 (Phase 3 pre-filter).

    The score is included in the compact payload as a hint for the LLM —
    a numeric signal it can weight when its judgment is uncertain."""
    # Stage 1: dedup + Smoke filter
    seen: set[str] = set()
    deduped: list[TestCase] = []
    for tc in analysis.exact_matched_tests + analysis.section_matched_tests:
        if tc.id in seen:
            continue
        if tc.sub_suite == "Smoke & Sanity":
            continue
        seen.add(tc.id)
        deduped.append(tc)

    # Stage 2: pre-filter by score. top_k = budget_hi * 4 gives the LLM ~4×
    # the ceiling to choose from — enough room for judgment, but far tighter
    # than the raw ~999-candidate pool.
    top_k = analysis.budget.final_hi * 4
    scored = pre_filter_candidates(deduped, analysis.scoring_context, top_k=top_k)

    out = [tc for tc, _ in scored]
    compact = []
    for tc, score in scored:
        row = {
            "id": tc.id,
            "title": tc.title,
            "section": tc.section_top,
            "suite": tc.sub_suite,
            "auto": tc.automation,
            "score": score,
        }
        # Only surface flag_state when it's a NON-DEFAULT classification.
        # Adding "unflagged" to every candidate would inflate the prompt
        # with noise (most candidates aren't gated).
        state = analysis.section_flag_state.get(tc.section_top)
        if state and state != "unflagged":
            row["flag_state"] = state
        compact.append(row)
    return out, compact


def _module_summary_for_prompt(analysis: Analysis) -> list[dict]:
    result: list[dict] = []
    for mc in sorted(analysis.module_changes.values(), key=lambda m: m.total_loc, reverse=True):
        row = {
            "module": mc.module,
            "loc": mc.total_loc,
            "files": [f.path for f in mc.files[:20]],   # cap to keep prompt reasonable
        }
        # Attach Nimbus flag state when this module is gated. Same rationale
        # as _candidates_for_prompt: only include when non-default, to avoid
        # inflating the prompt with "unflagged" noise on every module.
        flag_info = analysis.module_flag_states.get(mc.module)
        if flag_info:
            row["flag_state"] = flag_info["classification"]
            row["gating_flags"] = flag_info["flags"]
        result.append(row)
    return result


def _risks_for_prompt(analysis: Analysis) -> list[dict]:
    return [
        {"kind": r.kind, "severity": r.severity, "location": r.location, "detail": r.detail}
        for r in analysis.risks
    ]


def _drift_for_prompt(analysis: Analysis) -> list[dict]:
    return [{"kind": d.kind, "item": d.item, "detail": d.detail} for d in analysis.drift]


def _flag_catalogue_for_prompt(analysis: Analysis) -> list[dict]:
    """Compact summary of every Nimbus flag in the catalogue with its
    classification. Gives the LLM a GLOBAL view of flag state so it can reason
    about features whose code isn't directly attributed (attribution is
    grep-based and misses indirect gating). Empty when Nimbus analysis was
    skipped (--repo-path not provided).

    Kept minimal — accessor + classification only. Detailed per-channel state
    is available via `gating_flags` in module_changes when a module IS
    attributed.
    """
    from nimbus_flag_analyzer import classify_gating_flags
    if not analysis.flag_catalogue:
        return []
    return [
        {
            "accessor": accessor,
            "classification": classify_gating_flags([accessor], analysis.flag_catalogue),
            "state": flag.channel_state_summary(),
        }
        for accessor, flag in sorted(analysis.flag_catalogue.items())
    ]


def _deterministic_rerank(analysis: Analysis) -> list[RankedTest]:
    """Fallback when LLM unavailable. Rough heuristic ordering by sub_suite and
    automation status — not a substitute for the LLM, but better than nothing.
    Respects the release-specific budget ceiling."""
    candidates, _ = _candidates_for_prompt(analysis)
    # Smoke & Sanity is filtered out upstream in _candidates_for_prompt.
    suite_order = {"Functional": 0, "Special Case": 1}
    auto_order = {"Unsuitable": 0, "Untriaged": 1, "Suitable": 2, "Disabled": 3, "Completed": 4}
    candidates.sort(key=lambda tc: (suite_order.get(tc.sub_suite, 9), auto_order.get(tc.automation, 9)))
    out: list[RankedTest] = []
    for tc in candidates[:analysis.budget.final_hi]:
        # crude priority: manual-only = P1 (must be run by humans), automated = P2
        if tc.automation in ("Unsuitable", "Untriaged"):
            prio = "P1"
        else:
            prio = "P2"
        out.append(RankedTest(test_id=tc.id, priority=prio,
                              reason=f"section '{tc.section_top}' was matched against touched modules (deterministic fallback)"))
    return out


def llm_rerank(analysis: Analysis) -> tuple[list[RankedTest], str]:
    """Call the configured rerank model (LLM_MODEL_RERANK) to prioritize
    candidates. Returns (ranked, notes). Falls back to deterministic ranker
    on any failure. Uses analysis.budget to size the request."""
    if not _llm_available():
        return _deterministic_rerank(analysis), ""

    candidates, compact = _candidates_for_prompt(analysis)
    if not compact:
        return [], ""

    user_payload = {
        "release": {"from": analysis.from_tag, "to": analysis.to_tag},
        "module_changes": _module_summary_for_prompt(analysis),
        "risks": _risks_for_prompt(analysis),
        "drift": _drift_for_prompt(analysis),
        "flag_catalogue": _flag_catalogue_for_prompt(analysis),
        "candidates": compact,
    }
    user_text = (
        "Rank the candidates for this release. Use the rubric in the system prompt.\n\n"
        "## Release data (JSON)\n"
        f"```json\n{json.dumps(user_payload, separators=(',', ':'))}\n```"
    )

    # Model-specific parameter compatibility:
    #   - `effort` is a Sonnet/Opus feature; Haiku rejects it.
    #   - `temperature` is deprecated on Sonnet 5+ / Opus 4.8+ (determinism
    #     is controlled implicitly by effort/thinking on those models).
    # Structured output via json_schema is universal.
    model_lc = LLM_MODEL_RERANK.lower()
    output_config: dict = {"format": {"type": "json_schema", "schema": RERANK_SCHEMA}}
    if "haiku" not in model_lc:
        output_config["effort"] = os.environ.get("RECOMMEND_RERANK_EFFORT", "low")

    # Inject the release-specific budget into the prompt. `{budget_lo}` and
    # `{budget_hi}` placeholders come from Phase 2. Everything else in the
    # prompt is stable, so the cached prefix is still hit.
    prompt_rendered = SYSTEM_PROMPT_RERANK.format(
        budget_lo=analysis.budget.final_lo,
        budget_hi=analysis.budget.final_hi,
    )

    kwargs: dict = {
        "model": LLM_MODEL_RERANK,
        "max_tokens": 8000,
        "thinking": {"type": "disabled"},
        "output_config": output_config,
        "system": [{
            "type": "text",
            "text": prompt_rendered,
            "cache_control": {"type": "ephemeral"},
        }],
        "messages": [{"role": "user", "content": user_text}],
    }
    # Only set temperature on models that still accept it.
    if "sonnet-5" not in model_lc and "opus-4-8" not in model_lc:
        kwargs["temperature"] = 0

    try:
        # max_retries=6 so the SDK rides out 30k-TPM rolling-window resets (~60s).
        client = anthropic.Anthropic(max_retries=6)
        response = client.messages.create(**kwargs)
    except anthropic.APIStatusError as e:
        sys.stderr.write(f"[recommend] LLM rerank failed ({e.status_code}): {e.message}\n")
        return _deterministic_rerank(analysis), ""
    except Exception as e:
        sys.stderr.write(f"[recommend] LLM rerank failed: {e}\n")
        return _deterministic_rerank(analysis), ""

    sys.stderr.write(
        f"[recommend]   rerank usage: input={response.usage.input_tokens} "
        f"cache_write={response.usage.cache_creation_input_tokens} "
        f"cache_read={response.usage.cache_read_input_tokens} "
        f"output={response.usage.output_tokens}\n"
    )

    text = next((b.text for b in response.content if b.type == "text"), "")
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as e:
        sys.stderr.write(f"[recommend] LLM rerank returned non-JSON: {e}\n")
        return _deterministic_rerank(analysis), ""

    valid_ids = {tc.id for tc in candidates}
    ranked: list[RankedTest] = []
    for item in parsed.get("ranked_tests", []):
        tid = item.get("test_id")
        prio = item.get("priority")
        reason = item.get("reason", "")
        if tid in valid_ids and prio in ("P0", "P1", "P2"):
            ranked.append(RankedTest(test_id=tid, priority=prio, reason=reason))

    # Defensive cap based on the release-specific budget. The prompt asks for
    # at most `budget_hi`, but some models (notably Haiku 4.5) do not respect
    # the instruction reliably. Enforce it here so downstream consumers can
    # trust the length regardless of model.
    hard_cap = analysis.budget.final_hi
    if len(ranked) > hard_cap:
        sys.stderr.write(
            f"[recommend] rerank returned {len(ranked)} tests, over the budget "
            f"ceiling {hard_cap}. Trimming to first {hard_cap} (order preserved).\n"
        )
        ranked = ranked[:hard_cap]

    return ranked, parsed.get("notes", "")


def llm_synthesize(analysis: Analysis, ranked: list[RankedTest], notes: str, test_index: dict[str, TestCase]) -> str:
    """Call the configured synthesize model (LLM_MODEL_SYNTHESIZE) to write
    the full Markdown report. Falls back to deterministic Markdown if LLM is
    unavailable or fails."""
    if not _llm_available():
        return render_deterministic_report(analysis, [test_index[r.test_id] for r in ranked if r.test_id in test_index])

    # Enrich ranked tests with title/section/automation so the LLM has full context.
    ranked_payload = []
    for r in ranked:
        tc = test_index.get(r.test_id)
        if not tc:
            continue
        ranked_payload.append({
            "test_id": r.test_id,
            "priority": r.priority,
            "reason": r.reason,
            "title": tc.title,
            "section": tc.section_hierarchy,
            "sub_suite": tc.sub_suite,
            "automation": tc.automation,
        })

    user_payload = {
        "release": {"from": analysis.from_tag, "to": analysis.to_tag},
        "stats": {
            "prs_in_scope": len(analysis.prs),
            "prs_skipped": len(analysis.skipped_prs),
            "files_changed": len(analysis.file_changes),
            "total_loc": sum(f.additions + f.deletions for f in analysis.file_changes),
            "modules_touched": len(analysis.module_changes),
            "unclassified_files": len(analysis.unclassified_files),
        },
        "module_changes": _module_summary_for_prompt(analysis),
        "risks": _risks_for_prompt(analysis),
        "drift": _drift_for_prompt(analysis),
        "flag_catalogue": _flag_catalogue_for_prompt(analysis),
        "ranked_tests": ranked_payload,
        "rerank_notes": notes,
        "unclassified_files": [{"path": f.path, "additions": f.additions, "deletions": f.deletions}
                               for f in analysis.unclassified_files[:30]],
        "largest_prs": [
            {"number": p.number, "title": p.title, "author": p.author,
             "additions": p.additions, "deletions": p.deletions}
            for p in sorted(analysis.prs, key=lambda p: p.additions + p.deletions, reverse=True)[:10]
        ],
        "skipped_prs": [{"number": p.number, "title": p.title, "reason": reason}
                        for p, reason in analysis.skipped_prs],
    }

    user_text = (
        "Write the release QA report. Follow the structure exactly as specified in the system prompt.\n\n"
        "## Release data (JSON)\n"
        f"```json\n{json.dumps(user_payload, separators=(',', ':'))}\n```"
    )

    try:
        client = anthropic.Anthropic(max_retries=6)
        with client.messages.stream(
            model=LLM_MODEL_SYNTHESIZE,
            max_tokens=16000,
            thinking={"type": "disabled"},
            output_config={"effort": "medium"},
            system=[{
                "type": "text",
                "text": SYSTEM_PROMPT_SYNTHESIZE,
                "cache_control": {"type": "ephemeral"},
            }],
            messages=[{"role": "user", "content": user_text}],
        ) as stream:
            final = stream.get_final_message()
    except anthropic.APIStatusError as e:
        sys.stderr.write(f"[recommend] LLM synthesize failed ({e.status_code}): {e.message}\n")
        return render_deterministic_report(analysis, [test_index[r.test_id] for r in ranked if r.test_id in test_index])
    except Exception as e:
        sys.stderr.write(f"[recommend] LLM synthesize failed: {e}\n")
        return render_deterministic_report(analysis, [test_index[r.test_id] for r in ranked if r.test_id in test_index])

    sys.stderr.write(
        f"[recommend]   synth usage: input={final.usage.input_tokens} "
        f"cache_write={final.usage.cache_creation_input_tokens} "
        f"cache_read={final.usage.cache_read_input_tokens} "
        f"output={final.usage.output_tokens}\n"
    )

    body = next((b.text for b in final.content if b.type == "text"), "").strip()

    # Prepend the drift warning + budget line (both deterministic, must
    # always be visible regardless of LLM output).
    header_parts = [
        f"# Release Test Recommendation — {analysis.from_tag} → {analysis.to_tag}\n",
        f"_Generated by recommend.py — rerank: {LLM_MODEL_RERANK}, synthesize: {LLM_MODEL_SYNTHESIZE}._\n",
        f"_{analysis.budget.summary_line()}_\n",
    ]
    if analysis.drift:
        header_parts.append("## ⚠️ Mapping drift detected\n")
        for d in analysis.drift:
            header_parts.append(f"- **{d.kind}** — `{d.item}` — {d.detail}")
        header_parts.append("\n_Run `align.py` to curate the mapping YAML._\n")
    return "\n".join(header_parts) + "\n" + body


# =============================================================================
# Report writer (deterministic fallback)
# =============================================================================


def render_deterministic_report(analysis: Analysis, ranked_tests: list[TestCase]) -> str:
    lines: list[str] = []
    lines.append(f"# Release Test Recommendation — {analysis.from_tag} → {analysis.to_tag}\n")
    lines.append(f"_Generated by recommend.py — deterministic fallback (LLM unavailable or failed)._\n")
    lines.append(f"_{analysis.budget.summary_line()}_\n")

    # Drift warning at the top
    if analysis.drift:
        lines.append("## ⚠️ Mapping drift detected\n")
        for d in analysis.drift:
            lines.append(f"- **{d.kind}** — `{d.item}` — {d.detail}")
        lines.append("\n_Run `align.py` to curate. The report below still reflects all data available._\n")

    # Executive summary
    lines.append("## Executive summary\n")
    lines.append(f"- **PRs in scope**: {len(analysis.prs)} (after filtering {len(analysis.skipped_prs)} low-impact)")
    lines.append(f"- **Files changed**: {len(analysis.file_changes)}")
    total_loc = sum(f.additions + f.deletions for f in analysis.file_changes)
    lines.append(f"- **Total LOC changed**: {total_loc}")
    lines.append(f"- **Modules touched**: {len(analysis.module_changes)}")
    lines.append(f"- **Unclassified files**: {len(analysis.unclassified_files)}")
    lines.append(f"- **Risk signals**: {len(analysis.risks)} ({sum(1 for r in analysis.risks if r.severity=='high')} high)")
    lines.append(f"- **Candidate tests**: {len(ranked_tests)}")
    lines.append("")

    # Modules with most change
    lines.append("## Most-changed modules\n")
    top_modules = sorted(analysis.module_changes.values(), key=lambda m: m.total_loc, reverse=True)
    for mc in top_modules[:15]:
        lines.append(f"- `{mc.module}` — {len(mc.files)} files, {mc.total_loc} LOC")
    lines.append("")

    # Risks
    if analysis.risks:
        lines.append("## Risks\n")
        by_severity = {"high": [], "medium": [], "low": []}
        for r in analysis.risks:
            by_severity[r.severity].append(r)
        for sev in ("high", "medium", "low"):
            if by_severity[sev]:
                lines.append(f"### {sev.upper()}\n")
                for r in by_severity[sev]:
                    lines.append(f"- **{r.kind}** at `{r.location}`: {r.detail}")
                lines.append("")

    # Ranked tests
    lines.append("## Suggested tests to run\n")
    if not ranked_tests:
        lines.append("_No candidate tests matched. Check the drift section and/or the mapping YAML._\n")
    else:
        # group by section_top for readability
        by_section: dict[str, list[TestCase]] = defaultdict(list)
        for tc in ranked_tests:
            by_section[tc.section_top].append(tc)
        for sec, group in sorted(by_section.items(), key=lambda kv: -len(kv[1])):
            lines.append(f"### {sec} ({len(group)} tests)\n")
            for tc in group[:50]:
                lines.append(f"- `{tc.id}` — {tc.title} _(suite: {tc.sub_suite}; automation: {tc.automation})_")
            if len(group) > 50:
                lines.append(f"- _… and {len(group)-50} more in this section_")
            lines.append("")

    # Unclassified files (so they're not lost)
    if analysis.unclassified_files:
        lines.append("## Unclassified changes — no section mapping\n")
        lines.append("_These files were modified but no TestRail section maps to their module. They are NOT silently dropped._\n")
        for fc in analysis.unclassified_files[:30]:
            lines.append(f"- `{fc.path}` (+{fc.additions} / -{fc.deletions})")
        if len(analysis.unclassified_files) > 30:
            lines.append(f"- _… and {len(analysis.unclassified_files)-30} more_")
        lines.append("")

    # High-risk PRs
    high_risk_prs = sorted(analysis.prs, key=lambda p: p.additions + p.deletions, reverse=True)[:10]
    if high_risk_prs:
        lines.append("## Largest PRs (candidates for deeper review)\n")
        for p in high_risk_prs:
            lines.append(f"- #{p.number} (+{p.additions}/-{p.deletions}) — {p.title[:90]} _by @{p.author}_")
        lines.append("")

    # Skipped PRs (for transparency)
    if analysis.skipped_prs:
        lines.append("## Skipped (auto-classified low-impact)\n")
        for p, reason in analysis.skipped_prs:
            lines.append(f"- #{p.number} — {p.title[:80]} _({reason})_")
        lines.append("")

    lines.append("---")
    lines.append("_End of deterministic report. When the LLM is available, its rerank/synthesize output replaces the test-prioritization and narrative sections._")
    return "\n".join(lines)


# =============================================================================
# Main
# =============================================================================


def analyze_nimbus_flags(
    file_changes: list[FileChange],
    module_changes: dict[str, ModuleChange],
    mapping: dict,
    repo_path: Optional[Path],
) -> tuple[dict[str, FeatureFlag], dict[str, dict], dict[str, str]]:
    """Load the Nimbus flag catalogue and attribute the diff to gating flags.

    Returns (catalogue, module_flag_states, section_flag_state).

    - catalogue: {accessor: FeatureFlag}
    - module_flag_states: {module_path: {"classification": str, "flags": [flag_summary_dict]}}
    - section_flag_state: {section_top: classification}

    When repo_path is None (or the directory is missing) everything returns
    empty, and downstream code silently skips the flag-aware logic.
    """
    if repo_path is None:
        return {}, {}, {}

    nimbus_dir = repo_path / "firefox-ios" / "nimbus-features"
    catalogue = load_flag_catalogue(nimbus_dir)
    if not catalogue:
        return {}, {}, {}

    # Two-hop wrapper: most product code goes through NimbusFeatureFlagLayer
    # (`.isEnabled(.foo)`) rather than referencing `features.fooFeature` directly.
    # Parse the layer file once to build the {enum_case: accessor} map so
    # wrapper-style references are attributed too. Missing file is fine —
    # just means we fall back to direct attribution only.
    wrapper_path = repo_path / "firefox-ios" / "Client" / "Nimbus" / "NimbusFeatureFlagLayer.swift"
    case_to_accessor: dict[str, str] = {}
    if wrapper_path.is_file():
        try:
            case_to_accessor = parse_wrapper_mapping(wrapper_path.read_text(errors="replace"))
        except OSError:
            pass

    # Grep the current on-disk content of each changed .swift file. The
    # local checkout may be at a different SHA than the release tags, but
    # flag-reference patterns are slow-moving invariants — this is close
    # enough for MVP. Falls back to the truncated patch when the file is
    # not present locally (deleted, or checkout is behind).
    #
    # Test and noise paths are excluded: a test file may reference a flag
    # (to exercise both branches) without the flag actually gating any
    # product code. Attributing tests would mislabel test trees as gated
    # by unrelated flags.
    swift_pairs: list[tuple[str, str]] = []
    swift_changes: list[FileChange] = []
    for fc in file_changes:
        if not fc.path.endswith(".swift"):
            continue
        if is_test_path(fc.path) or is_noise_path(fc.path):
            continue
        local = repo_path / fc.path
        if local.is_file():
            try:
                content = local.read_text(errors="replace")
            except OSError:
                content = fc.patch
        else:
            content = fc.patch
        swift_pairs.append((fc.path, content))
        swift_changes.append(fc)

    attributions = attribute_files(swift_pairs, catalogue, case_to_accessor)

    # Aggregate file attributions up to the module level.
    known = known_modules_from_mapping(mapping)
    module_flag_accessors: dict[str, list[str]] = {}
    for fc, attr in zip(swift_changes, attributions):
        module = classify_file(fc.path, known)
        if module and attr.gating_flags:
            bucket = module_flag_accessors.setdefault(module, [])
            for acc in attr.gating_flags:
                if acc not in bucket:
                    bucket.append(acc)

    module_flag_states: dict[str, dict] = {}
    for module, accessors in module_flag_accessors.items():
        module_flag_states[module] = {
            "classification": classify_gating_flags(accessors, catalogue),
            "flags": [
                {
                    "accessor": acc,
                    "state": catalogue[acc].channel_state_summary(),
                }
                for acc in accessors if acc in catalogue
            ],
        }

    # Aggregate at the section level for candidate-side lookup.
    # A section's classification is the most-permissive across its touched
    # modules (matches classify_gating_flags's own aggregation rule).
    section_priority = {"unflagged": 0, "fully_dark": 1, "beta_only": 2, "enabled": 3}
    section_flag_state: dict[str, str] = {}
    for section_def in mapping.get("sections", []):
        section_name = section_def["name"]
        touched_modules = [
            m["path"] for m in section_def.get("modules", [])
            if m["path"] in module_changes
        ]
        # Only sections whose touched modules ALL have a Nimbus classification
        # get an aggregate state — otherwise we'd mislabel a section that has
        # both unflagged and dark modules as "dark".
        gated_states = [
            module_flag_states[m]["classification"]
            for m in touched_modules
            if m in module_flag_states
        ]
        if gated_states and len(gated_states) == len(touched_modules):
            section_flag_state[section_name] = max(gated_states, key=lambda s: section_priority[s])

    return catalogue, module_flag_states, section_flag_state


def run_pipeline(from_tag: str, to_tag: str, testrail_path: Path, mapping_path: Path, output_path: Path, verbose: bool = False, repo_path: Optional[Path] = None) -> None:
    def vlog(msg: str) -> None:
        if verbose:
            print(f"[recommend] {msg}", file=sys.stderr)

    vlog("loading mapping + testrail export …")
    mapping = load_mapping(mapping_path)
    tests = load_testrail(testrail_path)
    vlog(f"  {len(tests)} TestRail cases loaded, {len(mapping.get('sections', []))} sections in YAML")

    vlog(f"fetching diff {from_tag}...{to_tag} …")
    if repo_path:
        vlog(f"  using local git clone at {repo_path} (no 300-file cap)")
        file_changes, commits = fetch_compare_git(repo_path, from_tag, to_tag)
    else:
        file_changes, commits = fetch_compare(from_tag, to_tag)
    vlog(f"  {len(file_changes)} files, {len(commits)} commits")

    vlog("resolving PRs from commits …")
    prs = fetch_prs_for_commits(commits)
    vlog(f"  {len(prs)} unique PRs")

    vlog("filtering low-impact PRs …")
    kept_prs: list[PR] = []
    skipped: list[tuple[PR, str]] = []
    for p in prs:
        reason = is_low_impact_pr(p)
        if reason:
            skipped.append((p, reason))
        else:
            kept_prs.append(p)
    vlog(f"  kept {len(kept_prs)}, skipped {len(skipped)}")

    vlog("detecting drift …")
    drift = detect_drift(file_changes, tests, mapping)
    vlog(f"  {len(drift)} drift findings")

    vlog("computing risk heuristics …")
    risks = detect_risks(kept_prs, file_changes)
    vlog(f"  {len(risks)} risk signals ({sum(1 for r in risks if r.severity=='high')} high)")

    vlog("grouping files by module …")
    module_changes, unclassified = group_by_module(file_changes, kept_prs, mapping)
    vlog(f"  {len(module_changes)} modules touched, {len(unclassified)} unclassified files")

    vlog("matching tests (exact, by automated test name) …")
    exact = exact_match_by_test_file(file_changes, tests)
    vlog(f"  {len(exact)} exact matches")

    vlog("matching tests (by section) …")
    section_tests = section_match_tests(list(module_changes.keys()), mapping, tests)
    vlog(f"  {len(section_tests)} section-matched")

    vlog("computing test budget …")
    release_type = detect_release_type(from_tag, to_tag)
    # NOTE: max_pr_loc uses kept_prs (post low-impact filter), NOT the raw
    # `prs` list. Rationale: a 3000-LOC strings import or Mergify backport
    # shouldn't inflate the test budget just because it's large — those
    # changes are behaviour-neutral. Only PRs that survive the low-impact
    # filter count toward the "big PR" budget bump.
    signal = ReleaseSignal(
        total_loc=sum(f.additions + f.deletions for f in file_changes),
        max_pr_loc=max((p.additions + p.deletions for p in kept_prs), default=0),
        high_severity_risk_count=sum(1 for r in risks if r.severity == "high"),
    )
    budget = compute_test_budget(release_type, signal)
    vlog(f"  release_type={release_type}  budget={budget.final_lo}-{budget.final_hi}  "
         f"(bump=+{budget.bump}: {budget.bump_reasons or 'none'})")

    vlog("building scoring context (Phase 3 pre-filter) …")
    scoring_context = build_scoring_context(exact, module_changes, risks, mapping)
    vlog(f"  {len(scoring_context.high_loc_modules)} high-LOC modules, "
         f"{len(scoring_context.sections_with_risk)} sections with risk signal")

    vlog("analyzing Nimbus feature flags …")
    flag_catalogue, module_flag_states, section_flag_state = analyze_nimbus_flags(
        file_changes, module_changes, mapping, repo_path,
    )
    if flag_catalogue:
        counts = defaultdict(int)
        for state in section_flag_state.values():
            counts[state] += 1
        vlog(f"  {len(flag_catalogue)} flags in catalogue, "
             f"{len(module_flag_states)} modules gated, "
             f"section states: {dict(counts) or 'none'}")
    else:
        vlog("  skipped (--repo-path not provided or nimbus-features missing)")

    analysis = Analysis(
        from_tag=from_tag, to_tag=to_tag,
        prs=kept_prs, skipped_prs=skipped,
        file_changes=file_changes,
        module_changes=module_changes,
        risks=risks, drift=drift,
        exact_matched_tests=exact,
        section_matched_tests=section_tests,
        unclassified_files=unclassified,
        budget=budget,
        scoring_context=scoring_context,
        flag_catalogue=flag_catalogue,
        module_flag_states=module_flag_states,
        section_flag_state=section_flag_state,
    )

    # Preview the pre-filter effect for visibility in the log
    raw_pool_size = len({tc.id for tc in exact + section_tests if tc.sub_suite != "Smoke & Sanity"})
    top_k = budget.final_hi * 4
    vlog(f"pre-filter: {raw_pool_size} candidates → top {min(top_k, raw_pool_size)} (budget_hi × 4 = {top_k})")

    vlog("LLM rerank …")
    ranked, rerank_notes = llm_rerank(analysis)
    vlog(f"  {len(ranked)} tests ranked (P0={sum(1 for r in ranked if r.priority=='P0')}, "
         f"P1={sum(1 for r in ranked if r.priority=='P1')}, "
         f"P2={sum(1 for r in ranked if r.priority=='P2')})")

    vlog("LLM synthesize …")
    test_index = {tc.id: tc for tc in tests}
    report = llm_synthesize(analysis, ranked, rerank_notes, test_index)

    output_path.write_text(report)
    print(f"wrote {output_path}  ({len(report)} chars)")


def main() -> None:
    p = argparse.ArgumentParser(description="Firefox iOS Test Recommender")
    p.add_argument("--from", dest="from_tag", required=True, help="Previous release tag, e.g. firefox-v150.0")
    p.add_argument("--to", dest="to_tag", required=True, help="New release tag, e.g. firefox-v151.0")
    p.add_argument("--testrail", required=True, type=Path, help="Path to TestRail export .xlsx")
    p.add_argument("--mapping", required=True, type=Path, help="Path to section_to_module_mapping.yaml")
    p.add_argument("--output", type=Path, default=None, help="Output Markdown path (default: release_report_<to_tag>.md)")
    p.add_argument("--repo-path", type=Path, default=None,
                   help="Path to a local firefox-ios clone. When provided, the diff is computed from local git "
                        "instead of the GitHub compare API (removes the 300-file cap) and Nimbus feature-flag "
                        "analysis is enabled (fully_dark/beta_only/enabled per release channel).")
    p.add_argument("--verbose", "-v", action="store_true")
    args = p.parse_args()

    output = args.output or Path(f"release_report_{args.to_tag}.md")
    run_pipeline(args.from_tag, args.to_tag, args.testrail, args.mapping, output,
                 verbose=args.verbose, repo_path=args.repo_path)


if __name__ == "__main__":
    main()
