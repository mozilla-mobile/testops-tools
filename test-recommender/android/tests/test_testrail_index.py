"""Tests for android/testrail_index.py.

The xlsx reader is exercised against a workbook written on the fly, so the
suite does not depend on a checked-in TestRail export. The repo scanner runs
against a temporary directory holding synthetic Kotlin sources.
"""

from __future__ import annotations

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

import openpyxl

from android import testrail_index as ti


# =============================================================================
# Cell parsing
# =============================================================================


class TestParseDeclaredTests(unittest.TestCase):
    def test_plain_value(self):
        self.assertEqual(
            ti.parse_declared_tests("ui/OnboardingTest#verifyTheNotificationsOnboardingCardTest"),
            [("OnboardingTest", "verifyTheNotificationsOnboardingCardTest")],
        )

    def test_http_prefix_from_testrail_autolinking(self):
        self.assertEqual(
            ti.parse_declared_tests("http://ui/OnboardingTest#verifyTheStartSyncingOnboardingCardTest"),
            [("OnboardingTest", "verifyTheStartSyncingOnboardingCardTest")],
        )

    def test_class_name_need_not_end_in_test(self):
        """`MainMenuTestCompose` is a real class name; anchoring the pattern on
        a trailing `Test` silently dropped 52 cells of the current export."""
        self.assertEqual(
            ti.parse_declared_tests("ui/MainMenuTestCompose#verifyTheHistoryMenuOptionTest"),
            [("MainMenuTestCompose", "verifyTheHistoryMenuOptionTest")],
        )

    def test_strips_pasted_rich_text(self):
        cell = ("<span style='color: rgb(32, 32, 32); font-family: Arial'>"
                "ui/HomeScreenTest#homeScreenItemsTest</span>")
        self.assertEqual(
            ti.parse_declared_tests(cell),
            [("HomeScreenTest", "homeScreenItemsTest")],
        )

    def test_multiple_values_deduplicated(self):
        cell = "ui/ATest#one\nui/ATest#two\nui/ATest#one"
        self.assertEqual(
            ti.parse_declared_tests(cell), [("ATest", "one"), ("ATest", "two")]
        )

    def test_prose_yields_nothing(self):
        """Preconditions text pasted into the wrong column must not parse."""
        self.assertEqual(
            ti.parse_declared_tests(
                'Longfox is enabled in Secret Settings. The app is running.'),
            [],
        )

    def test_empty_and_none(self):
        self.assertEqual(ti.parse_declared_tests(None), [])
        self.assertEqual(ti.parse_declared_tests(""), [])


class TestSplitHierarchy(unittest.TestCase):
    def test_plain_hierarchy(self):
        self.assertEqual(
            ti.split_hierarchy("Menus > Settings menu > General > Homepage"),
            ["Menus", "Settings menu", "General", "Homepage"])

    def test_single_level(self):
        self.assertEqual(ti.split_hierarchy("Downloads"), ["Downloads"])

    def test_bracketed_separator_is_not_a_level(self):
        """`[Android > 14]` is a platform gate inside one section name, not a
        hierarchy boundary. Splitting naively truncated this to
        `Accounts & Sync > [Android`."""
        self.assertEqual(
            ti.split_hierarchy(
                "Accounts & Sync > [Android > 14] Send tabs through native share sheet"),
            ["Accounts & Sync", "[Android > 14] Send tabs through native share sheet"])

    def test_bracketed_separator_at_top_level(self):
        self.assertEqual(
            ti.split_hierarchy("[Android > 14] Native share sheet options"),
            ["[Android > 14] Native share sheet options"])

    def test_less_than_needs_no_special_handling(self):
        self.assertEqual(
            ti.split_hierarchy("Homescreen > [Android < 10] Messaging banner"),
            ["Homescreen", "[Android < 10] Messaging banner"])

    def test_empty_levels_are_dropped(self):
        self.assertEqual(ti.split_hierarchy("A >  > B"), ["A", "B"])


class TestSectionAt(unittest.TestCase):
    def _case(self, hierarchy):
        return ti.TestCase(
            case_id="1", title="t", section_hierarchy=hierarchy, priority="High",
            automation="Completed", automation_coverage="Full",
            case_type="Functional", labels="",
        )

    def test_depth_two_splits_the_menus_bucket(self):
        c = self._case("Menus > Settings menu > General > Homepage")
        self.assertEqual(c.section_at(1), "Menus")
        self.assertEqual(c.section_at(2), "Menus > Settings menu")

    def test_depth_beyond_hierarchy_returns_whole_path(self):
        c = self._case("Downloads")
        self.assertEqual(c.section_at(2), "Downloads")


class TestAutomationPartition(unittest.TestCase):
    def _case(self, automation):
        return ti.TestCase(
            case_id="1", title="t", section_hierarchy="Downloads", priority="High",
            automation=automation, automation_coverage="None",
            case_type="Functional", labels="",
        )

    def test_completed_and_disabled_count_as_automated(self):
        self.assertTrue(self._case("Completed").is_automated)
        self.assertTrue(self._case("Disabled").is_automated)

    def test_unsuitable_is_the_manual_only_pool(self):
        c = self._case("Unsuitable")
        self.assertTrue(c.is_manual_only)
        self.assertFalse(c.is_automated)

    def test_suitable_is_an_automation_gap(self):
        c = self._case("Suitable")
        self.assertTrue(c.is_automation_gap)
        self.assertFalse(c.is_manual_only)

    def test_untriaged_is_neither(self):
        c = self._case("Untriaged")
        self.assertFalse(c.is_automated)
        self.assertFalse(c.is_manual_only)
        self.assertFalse(c.is_automation_gap)


# =============================================================================
# Fixtures
# =============================================================================


HEADERS = [
    "ID", "Title", "Section Hierarchy", "Priority", "Automation",
    "Automation Coverage", "Type", "Labels", "Automated Test Name(s)",
]


def write_export(path: Path, rows: list[dict]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for row in rows:
        ws.append([row.get(h, "") for h in HEADERS])
    wb.save(str(path))


def write_test_file(root: Path, rel: str, body: str) -> None:
    target = root / ti.ANDROID_TEST_ROOT / rel
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(body, encoding="utf-8")


LINKED_TEST = """
package org.mozilla.fenix.ui

class BookmarksTest {
    // TestRail link: https://mozilla.testrail.io/index.php?/cases/view/2833690
    @SmokeTest
    @Test
    fun deleteBookmarkFoldersTest() {
        doSomething()
    }

    // TestRail link: https://mozilla.testrail.io/index.php?/cases/view/2833691
    @Test
    fun editBookmarkTest() {
    }

    private fun helper() {
    }
}
"""


# =============================================================================
# Export loading
# =============================================================================


class TestLoadExport(unittest.TestCase):
    def test_reads_rows_and_strips_case_prefix(self):
        with TemporaryDirectory() as tmp:
            p = Path(tmp) / "e.xlsx"
            write_export(p, [{
                "ID": "C2833690", "Title": "Delete bookmark folders",
                "Section Hierarchy": "Bookmarks", "Priority": "High",
                "Automation": "Completed", "Automation Coverage": "Full",
                "Type": "Functional",
                "Automated Test Name(s)": "ui/BookmarksTest#deleteBookmarkFoldersTest",
            }])
            cases = ti.load_export(p)
        self.assertIn("2833690", cases)
        self.assertEqual(cases["2833690"].declared_tests,
                         [("BookmarksTest", "deleteBookmarkFoldersTest")])

    def test_accepts_id_without_c_prefix(self):
        with TemporaryDirectory() as tmp:
            p = Path(tmp) / "e.xlsx"
            write_export(p, [{"ID": "2833690", "Title": "x",
                              "Section Hierarchy": "Bookmarks"}])
            self.assertIn("2833690", ti.load_export(p))

    def test_skips_rows_without_a_usable_id(self):
        with TemporaryDirectory() as tmp:
            p = Path(tmp) / "e.xlsx"
            write_export(p, [
                {"ID": "", "Title": "blank"},
                {"ID": "not-an-id", "Title": "junk"},
                {"ID": "C1", "Title": "ok", "Section Hierarchy": "Downloads"},
            ])
            self.assertEqual(list(ti.load_export(p)), ["1"])

    def test_missing_required_column_is_a_loud_error(self):
        with TemporaryDirectory() as tmp:
            p = Path(tmp) / "e.xlsx"
            wb = openpyxl.Workbook()
            wb.active.append(["ID", "Title"])       # no Section Hierarchy
            wb.save(str(p))
            with self.assertRaises(ValueError) as ctx:
                ti.load_export(p)
        self.assertIn("Section Hierarchy", str(ctx.exception))

    def test_column_order_does_not_matter(self):
        with TemporaryDirectory() as tmp:
            p = Path(tmp) / "e.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["Section Hierarchy", "Automation", "Title", "ID"])
            ws.append(["Downloads", "Unsuitable", "A manual case", "C7"])
            wb.save(str(p))
            cases = ti.load_export(p)
        self.assertEqual(cases["7"].automation, "Unsuitable")
        self.assertEqual(cases["7"].section_hierarchy, "Downloads")


# =============================================================================
# Repo scanning
# =============================================================================


class TestScanTestMethods(unittest.TestCase):
    def test_links_comment_to_the_following_test_method(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            write_test_file(root, "java/org/mozilla/fenix/ui/BookmarksTest.kt", LINKED_TEST)
            methods = ti.scan_test_methods(root)
        by_name = {m.method: m for m in methods}
        self.assertEqual(sorted(by_name), ["deleteBookmarkFoldersTest", "editBookmarkTest"])
        self.assertEqual(by_name["deleteBookmarkFoldersTest"].case_ids, ["2833690"])
        self.assertEqual(by_name["editBookmarkTest"].case_ids, ["2833691"])

    def test_private_helpers_are_not_tests(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            write_test_file(root, "java/org/mozilla/fenix/ui/BookmarksTest.kt", LINKED_TEST)
            methods = ti.scan_test_methods(root)
        self.assertNotIn("helper", [m.method for m in methods])

    def test_class_name_comes_from_the_filename(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            write_test_file(root, "java/org/mozilla/fenix/ui/BookmarksTest.kt", LINKED_TEST)
            methods = ti.scan_test_methods(root)
        self.assertEqual({m.class_name for m in methods}, {"BookmarksTest"})

    def test_paths_are_repo_relative(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            write_test_file(root, "java/org/mozilla/fenix/ui/BookmarksTest.kt", LINKED_TEST)
            methods = ti.scan_test_methods(root)
        self.assertTrue(methods[0].path.startswith(ti.ANDROID_TEST_ROOT))
        self.assertFalse(Path(methods[0].path).is_absolute())

    def test_file_without_tests_is_skipped(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            write_test_file(root, "java/org/mozilla/fenix/helpers/Helper.kt",
                            "class Helper { fun doThing() {} }")
            self.assertEqual(ti.scan_test_methods(root), [])

    def test_missing_test_root_returns_empty(self):
        with TemporaryDirectory() as tmp:
            self.assertEqual(ti.scan_test_methods(Path(tmp)), [])

    def test_malformed_declaration_does_not_leak_links_to_the_next_method(self):
        """A working tree containing `fun () {` moved case 111 onto the method
        below it. The links must be dropped and the anomaly reported."""
        body = """
class BrokenTest {
    // TestRail link: https://mozilla.testrail.io/index.php?/cases/view/111
    @Test
    fun () {
    }

    // TestRail link: https://mozilla.testrail.io/index.php?/cases/view/222
    @Test
    fun theNextOneTest() {
    }
}
"""
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            write_test_file(root, "java/org/mozilla/fenix/ui/BrokenTest.kt", body)
            malformed: list[tuple[str, str]] = []
            methods = ti.scan_test_methods(root, malformed=malformed)
        self.assertEqual([m.method for m in methods], ["theNextOneTest"])
        self.assertEqual(methods[0].case_ids, ["222"])
        self.assertEqual(len(malformed), 1)
        self.assertEqual(malformed[0][1], "fun () {")

    def test_malformed_declaration_surfaces_as_drift(self):
        body = """
class BrokenTest {
    // TestRail link: https://mozilla.testrail.io/index.php?/cases/view/111
    @Test
    fun () {}
}
"""
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            export = root / "e.xlsx"
            write_export(export, [{"ID": "C111", "Title": "x",
                                   "Section Hierarchy": "Bookmarks",
                                   "Automation": "Completed"}])
            write_test_file(root, "java/org/mozilla/fenix/ui/BrokenTest.kt", body)
            ix = ti.build_index(export, root)
        self.assertEqual(len(ix.drift.malformed_methods), 1)
        self.assertFalse(ix.drift.is_clean)
        self.assertNotIn("111", ix.case_to_files)

    def test_modifier_prefixed_helpers_still_do_not_consume_links(self):
        """`private fun` between a link and its @Test must not be mistaken for
        a malformed test declaration."""
        body = """
class OkTest {
    // TestRail link: https://mozilla.testrail.io/index.php?/cases/view/333
    @Test
    fun realTest() {}

    private fun helper() {}
}
"""
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            write_test_file(root, "java/org/mozilla/fenix/ui/OkTest.kt", body)
            malformed: list[tuple[str, str]] = []
            methods = ti.scan_test_methods(root, malformed=malformed)
        self.assertEqual(malformed, [])
        # `private fun` is not captured as a method at all — the modifier makes
        # KOTLIN_FUN_RE miss it — and must not be mistaken for a broken @Test.
        self.assertEqual([m.method for m in methods], ["realTest"])
        self.assertEqual(methods[0].case_ids, ["333"])

    def test_multiple_links_on_one_method(self):
        body = """
class MultiTest {
    // TestRail link: https://mozilla.testrail.io/index.php?/cases/view/111
    // TestRail link: https://mozilla.testrail.io/index.php?/cases/view/222
    @Test
    fun coversTwoCasesTest() {}
}
"""
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            write_test_file(root, "java/org/mozilla/fenix/ui/MultiTest.kt", body)
            methods = ti.scan_test_methods(root)
        self.assertEqual(methods[0].case_ids, ["111", "222"])


# =============================================================================
# Index construction
# =============================================================================


class TestBuildIndex(unittest.TestCase):
    def _fixture(self, tmp: str, rows: list[dict], test_body: str = LINKED_TEST):
        root = Path(tmp)
        export = root / "export.xlsx"
        write_export(export, rows)
        write_test_file(root, "java/org/mozilla/fenix/ui/BookmarksTest.kt", test_body)
        return ti.build_index(export, root)

    def test_both_sources_agree(self):
        rows = [{
            "ID": "C2833690", "Title": "Delete bookmark folders",
            "Section Hierarchy": "Bookmarks", "Automation": "Completed",
            "Automated Test Name(s)": "ui/BookmarksTest#deleteBookmarkFoldersTest",
        }]
        with TemporaryDirectory() as tmp:
            ix = self._fixture(tmp, rows)
        path = f"{ti.ANDROID_TEST_ROOT}/java/org/mozilla/fenix/ui/BookmarksTest.kt"
        self.assertEqual(ix.case_to_files["2833690"], {path})
        self.assertIn("2833690", ix.file_to_cases[path])
        self.assertEqual(ix.drift.stale_declared, [])

    def test_repo_only_link_still_maps(self):
        """A case linked from code but with an empty export column."""
        rows = [{"ID": "C2833690", "Title": "x", "Section Hierarchy": "Bookmarks",
                 "Automation": "Completed"}]
        with TemporaryDirectory() as tmp:
            ix = self._fixture(tmp, rows)
        self.assertIn("2833690", ix.case_to_files)

    def test_orphan_link_is_reported(self):
        """A `// TestRail link:` pointing at a case absent from the export."""
        rows = [{"ID": "C999", "Title": "x", "Section Hierarchy": "Bookmarks",
                 "Automation": "Completed"}]
        with TemporaryDirectory() as tmp:
            ix = self._fixture(tmp, rows)
        self.assertIn("2833690", ix.drift.orphan_links)
        self.assertIn("2833691", ix.drift.orphan_links)

    def test_stale_method_falls_back_to_the_class_file(self):
        """A renamed method must not lose the whole case mapping."""
        rows = [{
            "ID": "C500", "Title": "x", "Section Hierarchy": "Bookmarks",
            "Automation": "Completed",
            "Automated Test Name(s)": "ui/BookmarksTest#methodThatWasRenamed",
        }]
        with TemporaryDirectory() as tmp:
            ix = self._fixture(tmp, rows)
        path = f"{ti.ANDROID_TEST_ROOT}/java/org/mozilla/fenix/ui/BookmarksTest.kt"
        self.assertEqual(ix.case_to_files["500"], {path})
        self.assertEqual(ix.drift.stale_declared,
                         [("500", "BookmarksTest", "methodThatWasRenamed")])

    def test_vanished_class_yields_no_mapping(self):
        """`ComposeBookmarksTest` was renamed away; nothing to map to."""
        rows = [{
            "ID": "C600", "Title": "x", "Section Hierarchy": "Bookmarks",
            "Automation": "Completed",
            "Automated Test Name(s)": "ui/ComposeBookmarksTest#deleteBookmarkFoldersTest",
        }]
        with TemporaryDirectory() as tmp:
            ix = self._fixture(tmp, rows)
        self.assertNotIn("600", ix.case_to_files)
        self.assertEqual(ix.drift.stale_declared,
                         [("600", "ComposeBookmarksTest", "deleteBookmarkFoldersTest")])

    def test_unparseable_cell_is_reported(self):
        rows = [{
            "ID": "C700", "Title": "x", "Section Hierarchy": "Bookmarks",
            "Automation": "Suitable",
            "Automated Test Name(s)": "Longfox is enabled in Secret Settings.",
        }]
        with TemporaryDirectory() as tmp:
            ix = self._fixture(tmp, rows)
        self.assertEqual(len(ix.drift.unparseable_declared), 1)
        self.assertEqual(ix.drift.unparseable_declared[0][0], "700")

    def test_automation_status_mismatch_is_reported(self):
        """Code links the case, but TestRail says it is not automated."""
        rows = [{"ID": "C2833690", "Title": "x", "Section Hierarchy": "Bookmarks",
                 "Automation": "Unsuitable"}]
        with TemporaryDirectory() as tmp:
            ix = self._fixture(tmp, rows)
        self.assertIn(("2833690", "Unsuitable"), ix.drift.automation_status_mismatch)

    def test_clean_index_reports_no_drift(self):
        rows = [
            {"ID": "C2833690", "Title": "x", "Section Hierarchy": "Bookmarks",
             "Automation": "Completed",
             "Automated Test Name(s)": "ui/BookmarksTest#deleteBookmarkFoldersTest"},
            {"ID": "C2833691", "Title": "y", "Section Hierarchy": "Bookmarks",
             "Automation": "Completed",
             "Automated Test Name(s)": "ui/BookmarksTest#editBookmarkTest"},
        ]
        with TemporaryDirectory() as tmp:
            ix = self._fixture(tmp, rows)
        self.assertTrue(ix.drift.is_clean, vars(ix.drift))

    def test_cases_for_files_is_the_exact_match_signal(self):
        rows = [
            {"ID": "C2833690", "Title": "x", "Section Hierarchy": "Bookmarks",
             "Automation": "Completed"},
            {"ID": "C2833691", "Title": "y", "Section Hierarchy": "Bookmarks",
             "Automation": "Completed"},
        ]
        with TemporaryDirectory() as tmp:
            ix = self._fixture(tmp, rows)
        path = f"{ti.ANDROID_TEST_ROOT}/java/org/mozilla/fenix/ui/BookmarksTest.kt"
        self.assertEqual(ix.cases_for_files([path]), {"2833690", "2833691"})
        self.assertEqual(ix.cases_for_files(["some/unrelated/File.kt"]), set())

    def test_by_section_groups_at_requested_depth(self):
        rows = [
            {"ID": "C1", "Title": "a",
             "Section Hierarchy": "Menus > Settings menu > General",
             "Automation": "Unsuitable"},
            {"ID": "C2", "Title": "b",
             "Section Hierarchy": "Menus > Settings menu > Tabs",
             "Automation": "Unsuitable"},
            {"ID": "C3", "Title": "c",
             "Section Hierarchy": "Menus > Menu Redesign 2025",
             "Automation": "Unsuitable"},
        ]
        with TemporaryDirectory() as tmp:
            ix = self._fixture(tmp, rows)
        depth2 = ix.by_section(depth=2)
        self.assertEqual(len(depth2["Menus > Settings menu"]), 2)
        self.assertEqual(len(depth2["Menus > Menu Redesign 2025"]), 1)
        self.assertEqual(len(ix.by_section(depth=1)["Menus"]), 3)


if __name__ == "__main__":
    unittest.main()
